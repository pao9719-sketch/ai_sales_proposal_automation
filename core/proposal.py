"""공급 제안서(외부 제출용)와 내부 손익 검토서(내부용) 생성. 두 문서는 의도적으로 분리한다 (기획서 4.4).

외부 제출용에는 공급원가·운영비·매출총이익·영업기여이익·회수기간 등 내부 손익 정보를 넣지 않는다.
"""
from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .calc import investment_by_category, item_amount


def _won(v: float) -> str:
    return f"{v:,.0f}원"


def _table(doc: Document, header: list[str], rows: list[list[str]]):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    return t


def _title(doc: Document, text: str):
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_external_proposal(snap: dict, opinion: str, company: str = "MS가스") -> bytes:
    p, items = snap["project"], snap["items"]
    doc = Document()
    _title(doc, "산업가스 공급 제안서")
    doc.add_paragraph(f"수신: {p['customer']} 귀중").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"제안일: {date.today():%Y-%m-%d}  /  제안사: {company}").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. 고객사 및 공급 사업 개요", level=1)
    doc.add_paragraph(f"건명: {p['name']}")
    doc.add_paragraph(f"고객사: {p['customer']}")
    if p.get("site_note"):
        doc.add_paragraph(f"현장 개요: {p['site_note']}")

    doc.add_heading("2. 산업가스 공급 조건", level=1)
    _table(doc, ["항목", "내용"], [
        ["공급 가스", p["gas"] or "-"],
        ["예상 연간 사용량", f"{p['annual_volume']:,.0f} {p['unit']}"],
        ["공급단가", f"{p['supply_price']:,.0f}원 / {p['unit']}"],
        ["공급 기간", f"{p['contract_years']}년"],
    ])

    doc.add_heading("3. 저장설비 설치 및 배관공사 개요", level=1)
    cats = investment_by_category(items)
    if cats:
        _table(doc, ["구분", "공사 내용"], [[c, ", ".join(i["name"] for i in items
                                              if (i.get("category") or "기타") == c)] for c in cats])
    else:
        doc.add_paragraph("협의 후 확정")

    doc.add_heading("4. 예상 투자비 및 공급단가", level=1)
    if cats:
        _table(doc, ["구분", "예상 금액"], [[c, _won(v)] for c, v in cats.items()]
               + [["합계", _won(sum(cats.values()))]])
    doc.add_paragraph(f"공급단가: {p['supply_price']:,.0f}원 / {p['unit']}")

    doc.add_heading("5. 공급 조건 및 계약 관련 사항", level=1)
    doc.add_paragraph("세부 계약 조건(결제 조건, 설비 소유 및 반환, 최소 사용량 등)은 별도 협의 후 계약서에 명시합니다.")

    doc.add_heading("6. 검토 의견 및 추가 협의사항", level=1)
    doc.add_paragraph(opinion or "추가 협의사항은 담당자와 협의 바랍니다.")

    doc.add_paragraph()
    doc.add_paragraph("※ 본 문서는 초안이며, 담당자 검토 후 제출합니다.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_internal_review(snap: dict, issues: list[str], audit: list[dict]) -> bytes:
    p, items, r = snap["project"], snap["items"], snap["pnl"]
    doc = Document()
    _title(doc, "내부 손익 검토서 (대외비)")
    doc.add_paragraph(f"건명: {p['name']} / 고객사: {p['customer']}")
    doc.add_paragraph(f"승인자: {snap.get('approved_by', '-')}  /  승인일시: {snap.get('approved_at', '-')}")

    doc.add_heading("1. 손익 요약", level=1)
    pb = "회수 불가" if r["payback_years"] is None else f"{r['payback_years']:.1f}년"
    _table(doc, ["항목", "값"], [
        ["예상 매출액(연)", _won(r["revenue"])],
        ["예상 매출원가(연)", _won(r["cogs"])],
        ["예상 매출총이익(연)", f"{_won(r['gross_profit'])} ({r['gross_margin_pct']:.1f}%)"],
        ["추가 운영비(연)", _won(p["annual_opex"])],
        ["예상 영업기여이익(연)", f"{_won(r['contribution'])} ({r['contribution_margin_pct']:.1f}%)"],
        ["초기 투자비", _won(r["investment"])],
        ["투자비 회수기간", pb],
    ])

    doc.add_heading("2. 투자비 내역", level=1)
    _table(doc, ["구분", "항목", "수량", "단가", "금액", "출처"], [
        [i.get("category") or "기타", i["name"], f"{i['qty']:,.2f}", _won(i["unit_price"]),
         _won(item_amount(i["qty"], i["unit_price"])), i.get("source") or ""] for i in items])

    doc.add_heading("3. 승인 기준 검토", level=1)
    if issues:
        for s in issues:
            doc.add_paragraph(s, style="List Bullet")
    else:
        doc.add_paragraph("설정된 승인 기준을 모두 충족합니다.")

    doc.add_heading("4. 변경 이력", level=1)
    _table(doc, ["일시", "작성자", "작업", "내용"],
           [[a["ts"], a["actor"], a["action"], a["detail"]] for a in reversed(audit)])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
