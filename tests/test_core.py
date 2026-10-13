import io

import pandas as pd

from core.calc import (PnLInput, check_approval_criteria, compute_pnl, investment_by_category,
                       simulate, total_investment)
from core.excel_io import parse_quote
from core.masking import mask_text
from core.proposal import build_external_proposal


def base():
    return PnLInput(annual_volume=100_000, supply_price=500, unit_cost=300, annual_opex=5_000_000,
                    investment=60_000_000)


def test_pnl():
    r = compute_pnl(base())
    assert r.revenue == 50_000_000
    assert r.cogs == 30_000_000
    assert r.gross_profit == 20_000_000
    assert r.contribution == 15_000_000
    assert r.payback_years == 4.0
    assert round(r.gross_margin_pct, 1) == 40.0


def test_payback_none_when_not_profitable():
    r = compute_pnl(PnLInput(1000, 100, 100, 0, 1_000_000))
    assert r.payback_years is None
    assert check_approval_criteria(r, 5, 10)


def test_investment():
    items = [{"category": "배관공사", "qty": 10, "unit_price": 100_000},
             {"category": "배관공사", "qty": 1, "unit_price": 500_000},
             {"category": None, "qty": 2, "unit_price": 50_000}]
    assert total_investment(items) == 1_600_000
    assert investment_by_category(items) == {"배관공사": 1_500_000, "기타": 100_000}


def test_simulate_grid():
    rows = simulate(base(), [-10, 0, 10], [0, 10])
    assert len(rows) == 6
    zero = next(r for r in rows if r["공급단가 변동(%)"] == 0 and r["투자비 변동(%)"] == 0)
    assert zero["회수기간(년)"] == 4.0
    up = next(r for r in rows if r["공급단가 변동(%)"] == 10 and r["투자비 변동(%)"] == 0)
    assert up["회수기간(년)"] < 4.0


def test_masking():
    t = mask_text("(주)가나다 123-45-67890 010-1234-5678 a@b.co.kr", {"(주)가나다": "고객사"})
    assert "가나다" not in t and "123-45-67890" not in t and "010-1234-5678" not in t and "a@b.co.kr" not in t


def test_parse_quote():
    df = pd.DataFrame([["공사 견적서", None, None], ["품명", "수량", "금액"],
                       ["저장탱크", 1, 30_000_000], ["배관", 50, 5_000_000], ["합계", None, 35_000_000]])
    buf = io.BytesIO()
    df.to_excel(buf, header=False, index=False)
    out = parse_quote(buf.getvalue())
    assert list(out["name"]) == ["저장탱크", "배관"]
    assert total_investment(out.to_dict("records")) == 35_000_000


def test_external_proposal_has_no_internal_figures():
    from docx import Document

    snap = {"project": {"name": "N", "customer": "C", "gas": "LN2", "annual_volume": 100_000, "unit": "Nm3",
                        "supply_price": 500, "unit_cost": 123_456_789, "annual_opex": 987_654_321,
                        "contract_years": 5, "site_note": ""},
            "items": [{"category": "배관공사", "name": "배관", "qty": 1, "unit_price": 1_000_000, "source": ""}]}
    data = build_external_proposal(snap, "의견")
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "123,456,789" not in text and "987,654,321" not in text
    assert "회수기간" not in text and "영업기여이익" not in text
